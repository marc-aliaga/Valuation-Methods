"""
Build the final Excel workbook in Amazon-template style.
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_LINE_ORDER, OUTPUT_LABELS


def build_dataframe(normalized_data, years):
    """
    Build a DataFrame:
        rows = canonical line items (in OUTPUT_LINE_ORDER)
        cols = years
    Uses the 'latest' value found in each statement.
    """
    rows = []
    for canonical in OUTPUT_LINE_ORDER:
        label = OUTPUT_LABELS.get(canonical, canonical)
        row = {"Line Item": label}

        for year in years:
            value = None
            if year in normalized_data:
                for statement in ("income", "balance_sheet", "cashflow"):
                    stmt = normalized_data[year].get(statement, {})
                    if canonical in stmt and stmt[canonical]:
                        value = stmt[canonical][0]   # take first (latest) value
                        break
            row[str(year)] = value

        rows.append(row)

    return pd.DataFrame(rows)


def add_growth_columns(df, years):
    """Add YoY growth columns for each year after the first."""
    for i in range(1, len(years)):
        prev_year = str(years[i - 1])
        curr_year = str(years[i])
        col_name = f"Growth {curr_year}"
        df[col_name] = df.apply(
            lambda r: (r[curr_year] / r[prev_year] - 1)
            if r[prev_year] not in (None, 0) and r[curr_year] is not None
            else None,
            axis=1,
        )
    return df


def write_excel(df, output_path, company_name="Rosneft Oil Company",
                ticker="ROSN", currency="RUB", unit="millions"):
    """Write the DataFrame to Excel with formatting."""
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Financials", index=False, startrow=4)

    wb = load_workbook(output_path)
    ws = wb["Financials"]

    # ---------- Header block ----------
    ws["A1"] = company_name
    ws["A1"].font = Font(bold=True, size=14)

    ws["A2"] = f"{ticker}    Currency: {currency}    Unit: {unit}"
    ws["A2"].font = Font(italic=True, size=10)

    ws["A3"] = "Source: Summary Consolidated Financial Statements (IFRS)"
    ws["A3"].font = Font(italic=True, size=9, color="666666")

    # ---------- Style the table header ----------
    header_row = 5
    for col_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # ---------- Number formatting ----------
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_idx in range(header_row + 1, header_row + 1 + len(df)):
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            if col_idx > 1:  # numeric columns
                header = df.columns[col_idx - 1]
                if header.startswith("Growth"):
                    cell.number_format = "0.0%"
                else:
                    cell.number_format = "#,##0"

    # ---------- Column widths ----------
    ws.column_dimensions["A"].width = 55
    for col_idx in range(2, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

    wb.save(output_path)
    print(f"✅ Excel saved: {output_path}")